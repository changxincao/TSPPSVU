import csv
import math
from pathlib import Path


BASE = Path(r"analysis/巴西数据分析/新版_purchase时间")
INCREMENT_ROOT = BASE / "第一组_C_h0.5增量补充"
OLD_FINAL_ROOT = BASE / "第一组_最优参数提取结果_4方法"
OUT_ROOT = BASE / "第一组_最优参数提取结果_4方法_含C0.5增量"


def read_csv(path: Path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def read_changed_actual_rows(path: Path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))
    header = rows[0]
    out = []
    for raw in rows[1:]:
        if not raw:
            continue
        # CSAA rows align with header directly.
        if raw[2] == "CSAA":
            out.append(dict(zip(header, raw)))
            continue
        # RCSAA rows split the param field into two columns: "C_h=0.5" and " lambda=..."
        # Rebuild into the original 42-column shape by merging those two pieces.
        if raw[2] == "RCSAA" and len(raw) == len(header):
            fixed = [None] * len(header)
            fixed[0] = raw[0]
            fixed[1] = raw[1]
            fixed[2] = raw[2]
            fixed[3] = raw[3] + ", " + raw[4]
            for i in range(4, len(header) - 1):
                fixed[i] = raw[i + 1]
            fixed[-1] = ""
            out.append(dict(zip(header, fixed)))
            continue
        raise ValueError(f"Unexpected changed actual row shape: method={raw[2]!r}, len={len(raw)}")
    return out


def write_csv(path: Path, rows, fieldnames):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(rows)


def mean(xs):
    return sum(xs) / len(xs) if xs else float("nan")


def stdev(xs):
    if len(xs) <= 1:
        return 0.0
    m = mean(xs)
    return math.sqrt(sum((x - m) ** 2 for x in xs) / (len(xs) - 1))


def q(vals, p):
    if not vals:
        return float("nan")
    xs = sorted(vals)
    if len(xs) == 1:
        return xs[0]
    idx = p * (len(xs) - 1)
    lo = int(math.floor(idx))
    hi = int(math.ceil(idx))
    if lo == hi:
        return xs[lo]
    w = idx - lo
    return xs[lo] * (1 - w) + xs[hi] * w


def fmt_sel(v):
    if v is None or v == "":
        return ""
    x = float(v)
    if abs(x - round(x)) <= 1e-9:
        return str(int(round(x)))
    return f"{x:.10f}".rstrip("0").rstrip(".")


def infer_raw_cols(sample_row):
    keys = list(sample_row.keys())
    return {
        "trial_id": keys[5],
        "test_index": keys[6],
        "expected_obj": keys[16],
        "realized_obj": keys[17],
        "solve_time": keys[18],
        "selected_count": keys[19],
        "transport_cost": keys[20],
        "spot_cost": keys[21],
        "penalty_cost": keys[22],
        "ess": keys[25],
    }


def build_replacement_row(old_row, raw_row, sel, method_name, raw_cols):
    row = dict(old_row)
    row["final_method_name"] = method_name
    row["param_strategy"] = (
        "trial-wise best (k, C_h)"
        if method_name == "CSAA"
        else "trial-wise best (k, C_h, lambda)"
    )
    row["trialId"] = str(sel["trialId"])
    row["actual_test_period"] = str(int(raw_row[raw_cols["test_index"]]) + int(raw_row["k1"]))
    row["selected_k"] = fmt_sel(sel["bestK"])
    row["selected_C_h"] = fmt_sel(sel["bestC_h"])
    row["selected_lambda"] = "" if method_name == "CSAA" else fmt_sel(sel["bestLambda"])
    row["stage1_val_mean"] = sel["stage1ValMean"]
    row["stage2_val_mean"] = sel["stage2ValMean"]

    for key in old_row.keys():
        if key in raw_row:
            row[key] = raw_row[key]
    return row


def summarize_trials(rows, raw_cols):
    groups = {}
    for row in rows:
        groups.setdefault(row["final_method_name"], []).append(row)
    out = []
    for method in ["Mean", "SAA", "CSAA", "RCSAA"]:
        rs = groups[method]
        exp = [float(r[raw_cols["expected_obj"]]) for r in rs]
        rea = [float(r[raw_cols["realized_obj"]]) for r in rs]
        transport = [float(r[raw_cols["transport_cost"]]) for r in rs]
        spot = [float(r[raw_cols["spot_cost"]]) for r in rs]
        penalty = [float(r[raw_cols["penalty_cost"]]) for r in rs]
        solve_t = [float(r[raw_cols["solve_time"]]) for r in rs]
        selected = [float(r[raw_cols["selected_count"]]) for r in rs]
        ess = [float(r[raw_cols["ess"]]) for r in rs]
        out.append(
            {
                "final_method_name": method,
                "param_strategy": rs[0]["param_strategy"],
                "trial_count": str(len(rs)),
                "mean_expected_obj": mean(exp),
                "std_expected_obj": stdev(exp),
                "min_expected_obj": min(exp),
                "p20_expected_obj": q(exp, 0.2),
                "p50_expected_obj": q(exp, 0.5),
                "p75_expected_obj": q(exp, 0.75),
                "p80_expected_obj": q(exp, 0.8),
                "p95_expected_obj": q(exp, 0.95),
                "max_expected_obj": max(exp),
                "mean_realized_obj": mean(rea),
                "std_realized_obj": stdev(rea),
                "min_realized_obj": min(rea),
                "p20_realized_obj": q(rea, 0.2),
                "p50_realized_obj": q(rea, 0.5),
                "p75_realized_obj": q(rea, 0.75),
                "p80_realized_obj": q(rea, 0.8),
                "p95_realized_obj": q(rea, 0.95),
                "max_realized_obj": max(rea),
                "mean_transport_cost": mean(transport),
                "mean_spot_cost": mean(spot),
                "mean_penalty_cost": mean(penalty),
                "mean_solve_time_sec": mean(solve_t),
                "mean_selected_count": mean(selected),
                "mean_ess": mean(ess),
            }
        )
    return out


def summarize_increment(compare_rows):
    updated = [r for r in compare_rows if r["isStage1Updated"].lower() == "true"]
    unchanged = [r for r in compare_rows if r["isStage1Updated"].lower() != "true"]
    return [
        {
            "group": "all",
            "trial_count": len(compare_rows),
            "updated_count": len(updated),
            "unchanged_count": len(unchanged),
            "mean_old_stage1": mean([float(r["oldStage1Mean"]) for r in compare_rows]),
            "mean_new_stage1": mean([float(r["updatedStage1Mean"]) for r in compare_rows]),
            "mean_old_stage2": mean([float(r["oldStage2Mean"]) for r in compare_rows]),
            "mean_new_stage2": mean([float(r["updatedStage2Mean"]) for r in compare_rows]),
        },
        {
            "group": "updated_only",
            "trial_count": len(updated),
            "updated_count": len(updated),
            "unchanged_count": 0,
            "mean_old_stage1": mean([float(r["oldStage1Mean"]) for r in updated]),
            "mean_new_stage1": mean([float(r["updatedStage1Mean"]) for r in updated]),
            "mean_old_stage2": mean([float(r["oldStage2Mean"]) for r in updated]),
            "mean_new_stage2": mean([float(r["updatedStage2Mean"]) for r in updated]),
        },
    ]


def write_xlsx(path, trials, summary, compare_rows, increment_summary):
    try:
        from openpyxl import Workbook
    except Exception:
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "final_trials_4methods"
    ws.append(list(trials[0].keys()))
    for row in trials:
        ws.append([row[k] for k in trials[0].keys()])

    ws2 = wb.create_sheet("final_summary_4methods")
    ws2.append(list(summary[0].keys()))
    for row in summary:
        ws2.append([row[k] for k in summary[0].keys()])

    ws3 = wb.create_sheet("c05_trial_compare")
    ws3.append(list(compare_rows[0].keys()))
    for row in compare_rows:
        ws3.append([row[k] for k in compare_rows[0].keys()])

    ws4 = wb.create_sheet("c05_increment_summary")
    ws4.append(list(increment_summary[0].keys()))
    for row in increment_summary:
        ws4.append([row[k] for k in increment_summary[0].keys()])
    wb.save(path)


def main():
    OUT_ROOT.mkdir(parents=True, exist_ok=True)
    old_trials = read_csv(OLD_FINAL_ROOT / "final_trials_4methods.csv")
    updated_sel_rows = read_csv(INCREMENT_ROOT / "cv_selected_params_含C0.5增量_merged.csv")
    changed_raw_rows = read_changed_actual_rows(INCREMENT_ROOT / "changed_trials_actual_results_merged.csv")
    compare_rows = read_csv(INCREMENT_ROOT / "trial_compare_C0.5_vs_old_merged.csv")
    raw_cols = infer_raw_cols(changed_raw_rows[0])

    updated_sel = {int(r["trialId"]): r for r in updated_sel_rows}
    changed_raw = {(int(r[raw_cols["trial_id"]]), r["method_name"]): r for r in changed_raw_rows}
    changed_trial_ids = {
        tid for tid, r in updated_sel.items() if r["isUpdatedByC0.5"].lower() == "true"
    }

    out_trials = []
    for row in old_trials:
        trial_id = int(row["trialId"])
        if trial_id in changed_trial_ids and row["final_method_name"] in {"CSAA", "RCSAA"}:
            method = row["final_method_name"]
            out_trials.append(
                build_replacement_row(
                    row, changed_raw[(trial_id, method)], updated_sel[trial_id], method, raw_cols
                )
            )
        else:
            out_trials.append(row)

    final_summary = summarize_trials(out_trials, raw_cols)
    increment_summary = summarize_increment(compare_rows)

    write_csv(OUT_ROOT / "final_trials_4methods_含C0.5增量.csv", out_trials, list(out_trials[0].keys()))
    write_csv(
        OUT_ROOT / "final_summary_4methods_含C0.5增量.csv",
        final_summary,
        list(final_summary[0].keys()),
    )
    write_csv(
        INCREMENT_ROOT / "C0.5增量_trial_compare_汇总.csv",
        compare_rows,
        list(compare_rows[0].keys()),
    )
    write_csv(
        INCREMENT_ROOT / "C0.5增量_summary.csv",
        increment_summary,
        list(increment_summary[0].keys()),
    )
    write_xlsx(
        OUT_ROOT / "24line_10供应商_滚动CV选参提取结果_4methods_含C0.5增量.xlsx",
        out_trials,
        final_summary,
        compare_rows,
        increment_summary,
    )
    with (OUT_ROOT / "changed_trials.txt").open("w", encoding="utf-8") as f:
        for trial_id in sorted(changed_trial_ids):
            f.write(f"{trial_id}\n")

    print(f"changed_trials={len(changed_trial_ids)}")
    print(f"out_root={OUT_ROOT}")


if __name__ == "__main__":
    main()
