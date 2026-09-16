from __future__ import annotations

import csv
import math
from collections import Counter
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "analysis" / "巴西数据分析" / "新版_purchase时间" / "输出"
RAW_CV_ROOT = (
    BASE
    / "04_第一组主实验_滚动CV选最优参数并提取最终CSAA_RCSAA结果"
    / "01_原始滚动CV选参输出"
)
CHUNKS_ROOT = RAW_CV_ROOT / "chunks"
OUT_ROOT = (
    BASE
    / "04_第一组主实验_滚动CV选最优参数并提取最终CSAA_RCSAA结果"
    / "02_最终结果_4方法"
)
BASELINE_TRIAL_CSV = (
    BASE
    / "03_10供应商基础网格与提取"
    / "按purchase时间_raw网格实验_10供应商"
    / "全部实验配置_4方法_对齐51次样本外结果"
    / "全部实验配置_对齐后51次样本外_trial.csv"
)


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path: Path, rows: list[dict[str, object]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def mean(xs: list[float]) -> float:
    return sum(xs) / len(xs) if xs else float("nan")


def stdev(xs: list[float]) -> float:
    if len(xs) <= 1:
        return 0.0
    m = mean(xs)
    return math.sqrt(sum((x - m) ** 2 for x in xs) / (len(xs) - 1))


def quantile(vals: list[float], p: float) -> float:
    if not vals:
        return float("nan")
    xs = sorted(vals)
    if len(xs) == 1:
        return xs[0]
    idx = p * (len(xs) - 1)
    lo = int(math.floor(idx))
    hi = int(math.ceil(idx))
    if lo == hi:
        return xs[lo]
    w = idx - lo
    return xs[lo] * (1 - w) + xs[hi] * w


def merge_chunk_files(file_name: str) -> list[dict[str, str]]:
    merged: list[dict[str, str]] = []
    for chunk_dir in sorted(CHUNKS_ROOT.iterdir()):
        if not chunk_dir.is_dir():
            continue
        path = chunk_dir / file_name
        if path.exists():
            merged.extend(read_csv(path))
    return merged


def ensure_selection_complete(selection_rows: list[dict[str, str]], actual_rows: list[dict[str, str]]) -> None:
    trial_ids = sorted(int(r["trialId"]) for r in selection_rows)
    if trial_ids != list(range(51)):
        raise ValueError(f"selection trial ids incomplete: got {trial_ids[:5]} ... {trial_ids[-5:]}")

    actual_pairs = Counter((int(r["trialId"]), r["method_name"]) for r in actual_rows)
    expected_pairs = Counter((trial_id, method) for trial_id in range(51) for method in ("CSAA", "RCSAA"))
    if actual_pairs != expected_pairs:
        missing = [k for k, v in expected_pairs.items() if actual_pairs[k] != v]
        raise ValueError(f"actual rows incomplete or duplicated, mismatched pairs: {missing[:10]}")


def build_baseline_rows() -> list[dict[str, object]]:
    rows = read_csv(BASELINE_TRIAL_CSV)
    out: list[dict[str, object]] = []
    for method_name in ("Mean", "SAA"):
        method_rows = [r for r in rows if r["method_name"] == method_name and r["k1"] == "3"]
        if len(method_rows) != 51:
            raise ValueError(f"{method_name} baseline rows with k=3 expected 51, got {len(method_rows)}")
        method_rows.sort(key=lambda r: int(r["试验ID"]))
        for src in method_rows:
            out.append(
                {
                    "final_method_name": method_name,
                    "param_strategy": "fixed baseline from 03 aligned trial table",
                    "trialId": src["试验ID"],
                    "actual_test_period": src["actual_test_period"],
                    "selected_k": "3",
                    "selected_C_h": "",
                    "selected_lambda": "",
                    "stage1_val_mean": "",
                    "stage2_val_mean": "",
                    "expected_obj": src["期望目标值"],
                    "realized_obj": src["实现目标值"],
                    "solve_time_sec": src["求解时间秒"],
                    "selected_count": src["选中供应商数量"],
                    "oos_transport_cost": src["样本外运输成本"],
                    "oos_spot_cost": src["样本外现货成本"],
                    "oos_penalty_cost": src["样本外罚金成本"],
                    "ESS": src["有效样本量"],
                    "top1W": src["最大权重"],
                    "top5Wsum": src["前5权重和"],
                    "source": "03_全部实验配置_对齐后51次样本外_trial",
                }
            )
    return out


def build_selected_rows(
    selection_rows: list[dict[str, str]], actual_rows: list[dict[str, str]]
) -> tuple[list[dict[str, object]], list[dict[str, object]]]:
    actual_index = {(int(r["trialId"]), r["method_name"]): r for r in actual_rows}

    final_rows: list[dict[str, object]] = []
    param_rows: list[dict[str, object]] = []
    for sel in sorted(selection_rows, key=lambda r: int(r["trialId"])):
        trial_id = int(sel["trialId"])
        csaa = actual_index[(trial_id, "CSAA")]
        rcsaa = actual_index[(trial_id, "RCSAA")]
        param_rows.append(
            {
                "trialId": sel["trialId"],
                "testPeriodIdx": sel["testPeriodIdx"],
                "actual_test_period": csaa["actual_test_period"],
                "Mean_k": "3",
                "SAA_k": "3",
                "CSAA_best_k": sel["bestK"],
                "CSAA_best_C_h": sel["bestC_h"],
                "RCSAA_best_k": sel["bestK"],
                "RCSAA_best_C_h": sel["bestC_h"],
                "RCSAA_best_lambda": sel["bestLambda"],
                "stage1_val_mean": sel["stage1ValMean"],
                "stage2_val_mean": sel["stage2ValMean"],
            }
        )
        for method_name, src in (("CSAA", csaa), ("RCSAA", rcsaa)):
            final_rows.append(
                {
                    "final_method_name": method_name,
                    "param_strategy": "trial-wise rolling CV best parameter",
                    "trialId": sel["trialId"],
                    "actual_test_period": src["actual_test_period"],
                    "selected_k": sel["bestK"],
                    "selected_C_h": sel["bestC_h"],
                    "selected_lambda": "" if method_name == "CSAA" else sel["bestLambda"],
                    "stage1_val_mean": sel["stage1ValMean"],
                    "stage2_val_mean": sel["stage2ValMean"],
                    "expected_obj": src["expected_obj"],
                    "realized_obj": src["realized_obj"],
                    "solve_time_sec": src["solve_time_sec"],
                    "selected_count": src["selected_count"],
                    "oos_transport_cost": src["oos_transport_cost"],
                    "oos_spot_cost": src["oos_spot_cost"],
                    "oos_penalty_cost": src["oos_penalty_cost"],
                    "ESS": src["ESS"],
                    "top1W": src["top1W"],
                    "top5Wsum": src["top5Wsum"],
                    "source": "04_rolling_CV_selected_actual_trials",
                }
            )
    return final_rows, param_rows


def summarize_trials(rows: list[dict[str, object]]) -> list[dict[str, object]]:
    groups: dict[str, list[dict[str, object]]] = {}
    for row in rows:
        groups.setdefault(str(row["final_method_name"]), []).append(row)

    summary_rows: list[dict[str, object]] = []
    method_order = ["Mean", "SAA", "CSAA", "RCSAA"]
    for method in method_order:
        rs = groups[method]
        exp = [float(r["expected_obj"]) for r in rs]
        rea = [float(r["realized_obj"]) for r in rs]
        transport = [float(r["oos_transport_cost"]) for r in rs]
        spot = [float(r["oos_spot_cost"]) for r in rs]
        penalty = [float(r["oos_penalty_cost"]) for r in rs]
        solve_t = [float(r["solve_time_sec"]) for r in rs]
        selected = [float(r["selected_count"]) for r in rs]
        ess = [float(r["ESS"]) for r in rs]
        summary_rows.append(
            {
                "final_method_name": method,
                "param_strategy": rs[0]["param_strategy"],
                "trial_count": len(rs),
                "mean_expected_obj": mean(exp),
                "std_expected_obj": stdev(exp),
                "min_expected_obj": min(exp),
                "p20_expected_obj": quantile(exp, 0.20),
                "p50_expected_obj": quantile(exp, 0.50),
                "p75_expected_obj": quantile(exp, 0.75),
                "p80_expected_obj": quantile(exp, 0.80),
                "p95_expected_obj": quantile(exp, 0.95),
                "max_expected_obj": max(exp),
                "mean_realized_obj": mean(rea),
                "std_realized_obj": stdev(rea),
                "min_realized_obj": min(rea),
                "p20_realized_obj": quantile(rea, 0.20),
                "p50_realized_obj": quantile(rea, 0.50),
                "p75_realized_obj": quantile(rea, 0.75),
                "p80_realized_obj": quantile(rea, 0.80),
                "p95_realized_obj": quantile(rea, 0.95),
                "max_realized_obj": max(rea),
                "mean_transport_cost": mean(transport),
                "mean_spot_cost": mean(spot),
                "mean_penalty_cost": mean(penalty),
                "mean_solve_time_sec": mean(solve_t),
                "mean_selected_count": mean(selected),
                "mean_ess": mean(ess),
            }
        )
    return summary_rows


def maybe_write_xlsx(
    path: Path,
    final_trials: list[dict[str, object]],
    final_summary: list[dict[str, object]],
    best_params: list[dict[str, object]],
) -> None:
    try:
        from openpyxl import Workbook
    except Exception:
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "final_trials_4methods"
    ws.append(list(final_trials[0].keys()))
    for row in final_trials:
        ws.append([row[k] for k in final_trials[0].keys()])

    ws2 = wb.create_sheet("final_summary_4methods")
    ws2.append(list(final_summary[0].keys()))
    for row in final_summary:
        ws2.append([row[k] for k in final_summary[0].keys()])

    ws3 = wb.create_sheet("best_params_by_trial")
    ws3.append(list(best_params[0].keys()))
    for row in best_params:
        ws3.append([row[k] for k in best_params[0].keys()])

    wb.save(path)


def main() -> None:
    OUT_ROOT.mkdir(parents=True, exist_ok=True)

    selection_rows = merge_chunk_files("cv_selected_params.csv")
    actual_rows = merge_chunk_files("cv_selected_actual_trials.csv")
    stage1_rows = merge_chunk_files("cv_stage1_k_c_candidates.csv")
    stage2_rows = merge_chunk_files("cv_stage2_lambda_candidates.csv")
    ensure_selection_complete(selection_rows, actual_rows)

    selection_rows.sort(key=lambda r: int(r["trialId"]))
    actual_rows.sort(key=lambda r: (int(r["trialId"]), r["method_name"]))
    stage1_rows.sort(key=lambda r: (int(r["trialId"]), int(r["k"]), float(r["C_h"])))
    stage2_rows.sort(key=lambda r: (int(r["trialId"]), int(r["k"]), float(r["C_h"]), float(r["lambda"])))

    if selection_rows:
        write_csv(RAW_CV_ROOT / "cv_selected_params.csv", selection_rows, list(selection_rows[0].keys()))
    if actual_rows:
        write_csv(RAW_CV_ROOT / "cv_selected_actual_trials.csv", actual_rows, list(actual_rows[0].keys()))
    if stage1_rows:
        write_csv(RAW_CV_ROOT / "cv_stage1_k_c_candidates.csv", stage1_rows, list(stage1_rows[0].keys()))
    if stage2_rows:
        write_csv(RAW_CV_ROOT / "cv_stage2_lambda_candidates.csv", stage2_rows, list(stage2_rows[0].keys()))

    baseline_rows = build_baseline_rows()
    selected_final_rows, best_param_rows = build_selected_rows(selection_rows, actual_rows)

    final_trials = baseline_rows + selected_final_rows
    method_rank = {"Mean": 0, "SAA": 1, "CSAA": 2, "RCSAA": 3}
    final_trials.sort(key=lambda r: (int(str(r["trialId"])), method_rank[str(r["final_method_name"])]))
    final_summary = summarize_trials(final_trials)

    write_csv(OUT_ROOT / "最终逐trial结果_4方法.csv", final_trials, list(final_trials[0].keys()))
    write_csv(OUT_ROOT / "最终汇总结果_4方法.csv", final_summary, list(final_summary[0].keys()))
    write_csv(OUT_ROOT / "每个trial最优参数汇总_4方法.csv", best_param_rows, list(best_param_rows[0].keys()))
    maybe_write_xlsx(
        OUT_ROOT / "24line_10供应商_滚动CV选参提取结果_4方法.xlsx",
        final_trials,
        final_summary,
        best_param_rows,
    )

    (OUT_ROOT / "说明.txt").write_text(
        "第一组主实验正式结果。\n"
        "CSAA/RCSAA 来自本次新版 purchase 数据滚动 CV 重跑后的 trial-wise 最优参数及最终样本外结果。\n"
        "Mean/SAA 来自 03 目录中对齐后 51 次样本外 trial 表的 k=3 基线结果。\n"
        "本次 C_h 候选包含 0.1、0.5、1、3、5、10、30、50、100。\n",
        encoding="utf-8",
    )

    print(f"merged selection rows={len(selection_rows)}")
    print(f"merged actual rows={len(actual_rows)}")
    print(f"final_trials rows={len(final_trials)}")
    print(f"final_summary rows={len(final_summary)}")
    print(f"best_param rows={len(best_param_rows)}")
    print(f"out_root={OUT_ROOT}")


if __name__ == "__main__":
    main()
